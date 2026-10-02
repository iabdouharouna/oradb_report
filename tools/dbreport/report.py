"""Generation des rapports (Markdown, HTML, CSV, Excel) a partir d'un snapshot."""

from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Any, Iterable

from jinja2 import Environment, FileSystemLoader, select_autoescape

from .model import DatabaseSnapshot

try:  # openpyxl n'est requis que pour le format Excel.
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    _HAS_OPENPYXL = True
except ImportError:  # pragma: no cover - selon l'environnement
    Workbook = None  # type: ignore[assignment]
    Alignment = Border = Font = PatternFill = Side = None  # type: ignore[assignment]
    get_column_letter = None  # type: ignore[assignment]
    _HAS_OPENPYXL = False

_INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
_MAX_SHEET_NAME = 31

MAX_MD_ROWS = 200
TEMPLATE_DIR = Path(__file__).resolve().parent / "templates"


def summary_row(snapshot: DatabaseSnapshot) -> dict[str, Any]:
    """Ligne de synthese du parc pour une base."""

    dg_sync = snapshot.fact("dataguard", "dg_sync_status")
    dg_mode = snapshot.fact("dataguard", "dg_protection_mode")
    dg_lag = snapshot.fact("dataguard", "dg_apply_lag") or snapshot.fact("backup", "dg_apply_lag")

    return {
        "Base": snapshot.name,
        "Env": snapshot.environment,
        "Version": snapshot.fact("identity", "version") or "",
        "Role": snapshot.fact("identity", "database_role") or "",
        "Taille (Mo)": snapshot.fact("capacity", "database_mb"),
        "TS max %": snapshot.fact("capacity", "tablespace_max_pct"),
        "Invalides": snapshot.fact("objects", "invalid_total"),
        "Verrouilles": snapshot.fact("security", "account_locked"),
        "Sauv. (j)": snapshot.fact("backup", "last_backup_age_days"),
        "DG Sync": dg_sync or dg_lag or "",
        "DG Mode": dg_mode or "",
        "Statut": "OK" if snapshot.ok else "ERREUR",
    }


def summary_rows(snapshots: Iterable[DatabaseSnapshot]) -> list[dict[str, Any]]:
    return [summary_row(snapshot) for snapshot in snapshots]


def _number(value: Any) -> float | None:
    return value if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def kpis(snapshots: list[DatabaseSnapshot]) -> dict[str, Any]:
    """Indicateurs cles du parc pour le tableau de bord HTML."""

    total_mb = 0.0
    max_ts = 0.0
    invalid = 0
    locked = 0
    errors = 0
    backups: list[float] = []
    dg_standbys = 0
    dg_synced = 0
    dg_gaps = 0
    for snapshot in snapshots:
        if not snapshot.ok:
            errors += 1
        mb = _number(snapshot.fact("capacity", "database_mb"))
        if mb is not None:
            total_mb += mb
        ts = _number(snapshot.fact("capacity", "tablespace_max_pct"))
        if ts is not None:
            max_ts = max(max_ts, ts)
        inv = _number(snapshot.fact("objects", "invalid_total"))
        if inv is not None:
            invalid += int(inv)
        lk = _number(snapshot.fact("security", "account_locked"))
        if lk is not None:
            locked += int(lk)
        age = _number(snapshot.fact("backup", "last_backup_age_days"))
        if age is not None:
            backups.append(age)
        dg_standbys += _number(snapshot.fact("dataguard", "dg_standby_count")) or 0
        dg_synced += _number(snapshot.fact("dataguard", "dg_synced_count")) or 0
        dg_gaps += _number(snapshot.fact("dataguard", "dg_gap_count")) or 0
    return {
        "bases": len(snapshots),
        "ok": len(snapshots) - errors,
        "errors": errors,
        "total_gb": round(total_mb / 1024, 1),
        "max_tablespace_pct": round(max_ts, 1),
        "invalid_total": invalid,
        "locked_total": locked,
        "oldest_backup_days": round(max(backups), 1) if backups else None,
        "dg_standby_total": int(dg_standbys),
        "dg_synced_total": int(dg_synced),
        "dg_gap_total": int(dg_gaps),
    }


def _tabs(snapshots: list[DatabaseSnapshot]) -> list[dict[str, Any]]:
    """Onglets du rapport HTML (un identifiant stable par base)."""

    tabs: list[dict[str, Any]] = []
    for index, snapshot in enumerate(snapshots):
        tabs.append(
            {
                "id": f"db-{index}",
                "name": snapshot.name,
                "environment": snapshot.environment,
                "ok": snapshot.ok,
            }
        )
    return tabs


def _md_table(rows: list[dict[str, Any]], max_rows: int = MAX_MD_ROWS) -> str:
    if not rows:
        return "_Aucune ligne._\n"
    columns = list(rows[0].keys())
    header = "| " + " | ".join(columns) + " |"
    separator = "| " + " | ".join("---" for _ in columns) + " |"
    lines = [header, separator]
    for row in rows[:max_rows]:
        cells = []
        for column in columns:
            value = row.get(column)
            cells.append("" if value is None else str(value).replace("|", "\\|"))
        lines.append("| " + " | ".join(cells) + " |")
    if len(rows) > max_rows:
        lines.append(f"\n_(tronque : {max_rows} lignes sur {len(rows)})._")
    return "\n".join(lines) + "\n"


def _facts_block(facts: dict[str, Any]) -> str:
    if not facts:
        return ""
    return "".join(f"- **{key}** : {value}\n" for key, value in facts.items())


def render_markdown(payload: dict[str, Any], snapshots: list[DatabaseSnapshot]) -> str:
    lines: list[str] = []
    lines.append("# Rapport standard du parc Oracle")
    lines.append("")
    lines.append(f"- **Genere le** : {payload.get('generated_at', '')}")
    lines.append(f"- **Inventaire** : {payload.get('inventory', '')}")
    lines.append(f"- **Bases** : {len(snapshots)}")
    lines.append("")
    lines.append("## Synthese du parc")
    lines.append("")
    lines.append(_md_table(summary_rows(snapshots), max_rows=500))

    for snapshot in snapshots:
        lines.append(f"## {snapshot.name}")
        lines.append("")
        lines.append(
            f"_Environnement : {snapshot.environment or '-'} | "
            f"Service : {snapshot.service or '-'} | "
            f"DSN : {snapshot.dsn} | Collecte : {snapshot.collected_at}_"
        )
        lines.append("")
        if not snapshot.ok:
            lines.append(f"> **ERREUR DE COLLECTE** : {snapshot.error}")
            lines.append("")
        for section in snapshot.sections:
            lines.append(f"### {section.section} — {section.status}")
            lines.append("")
            if section.facts:
                lines.append(_facts_block(section.facts))
                lines.append("")
            for table, rows in section.tables.items():
                lines.append(f"#### {table}")
                lines.append("")
                lines.append(_md_table(rows))
                lines.append("")
            if section.messages:
                lines.append("**Messages / avertissements :**")
                lines.append("")
                lines.append("".join(f"- {message}\n" for message in section.messages))
                lines.append("")
    return "\n".join(lines) + "\n"


def render_html(payload: dict[str, Any], snapshots: list[DatabaseSnapshot]) -> str:
    environment = Environment(
        loader=FileSystemLoader(str(TEMPLATE_DIR)),
        autoescape=select_autoescape(["html", "xml"]),
        trim_blocks=True,
        lstrip_blocks=True,
    )
    template = environment.get_template("report.html.j2")
    return template.render(
        meta=payload,
        snapshots=snapshots,
        summary=summary_rows(snapshots),
        kpis=kpis(snapshots),
        tabs=_tabs(snapshots),
    )


def _csv_columns(rows: list[dict[str, Any]]) -> list[str]:
    columns: list[str] = []
    for row in rows:
        for key in row:
            if key not in columns:
                columns.append(key)
    return columns


def write_csv(snapshots: list[DatabaseSnapshot], directory: Path) -> list[Path]:
    """Ecrit un fichier CSV par base / section / tableau, plus les faits."""

    written: list[Path] = []
    csv_dir = directory / "csv"
    csv_dir.mkdir(parents=True, exist_ok=True)

    for snapshot in snapshots:
        for section in snapshot.sections:
            for table, rows in section.tables.items():
                if not rows:
                    continue
                path = csv_dir / f"{snapshot.name}__{section.section}__{table}.csv"
                with path.open("w", encoding="utf-8", newline="") as handle:
                    writer = csv.DictWriter(handle, fieldnames=_csv_columns(rows))
                    writer.writeheader()
                    writer.writerows(rows)
                written.append(path)

        facts = [
            (section.section, key, value)
            for section in snapshot.sections
            for key, value in section.facts.items()
        ]
        if facts:
            path = csv_dir / f"{snapshot.name}__facts.csv"
            with path.open("w", encoding="utf-8", newline="") as handle:
                writer = csv.writer(handle)
                writer.writerow(["section", "key", "value"])
                writer.writerows(facts)
            written.append(path)

    return written


def _sheet_safe_name(name: str, used: set[str]) -> str:
    base = _INVALID_SHEET_CHARS.sub("_", name).strip() or "BASE"
    base = base[:_MAX_SHEET_NAME]
    candidate = base
    index = 1
    while candidate.lower() in used:
        suffix = f"_{index}"
        candidate = base[: _MAX_SHEET_NAME - len(suffix)] + suffix
        index += 1
    used.add(candidate.lower())
    return candidate


# Palette moderne du classeur Excel.
_INK = "1E293B"
_SLATE = "334155"
_BAND = "F1F5F9"
_ACCENT = "2563EB"
_OK = "16A34A"
_WARN = "D97706"
_ERR = "DC2626"
_MUTED = "94A3B8"


def _thin_border():
    edge = Side(style="thin", color="E2E8F0")
    return Border(left=edge, right=edge, top=edge, bottom=edge)


def _band(sheet, row: int, ncols: int, text: str) -> int:
    """Bandeau de section (fusionne et colore)."""

    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(ncols, 2))
    cell = sheet.cell(row=row, column=1, value=text)
    cell.font = Font(bold=True, color="FFFFFF", size=11)
    cell.fill = PatternFill("solid", fgColor=_SLATE)
    cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[row].height = 19
    return row + 1


def _threshold_color(column: str, value: Any) -> str | None:
    """Couleur d'alerte selon le nom de colonne et la valeur."""

    name = column.lower()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if "pct" in name:
            if value >= 90:
                return _ERR
            if value >= 80:
                return _WARN
        if "age" in name and value >= 7:
            return _WARN
    if isinstance(value, str):
        upper = value.upper()
        if upper in ("FAILED", "INVALID", "ERROR", "STOPPED", "BLOCKED"):
            return _ERR
        if upper in ("COMPLETED", "VALID", "OK", "SUCCESS", "OPEN"):
            return _OK
        if "PENDING" in upper or "PARTIAL" in upper or "PROGRESS" in upper:
            return _WARN
    return None


def _table_header(sheet, row: int, columns: list[str]) -> int:
    for col, name in enumerate(columns, start=1):
        cell = sheet.cell(row=row, column=col, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=_INK)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = _thin_border()
    sheet.row_dimensions[row].height = 17
    return row + 1


def _table_rows(sheet, row: int, rows: list[dict[str, Any]]) -> int:
    columns = _csv_columns(rows)
    row = _table_header(sheet, row, columns)
    for position, item in enumerate(rows):
        for col, name in enumerate(columns, start=1):
            value = item.get(name)
            cell = sheet.cell(row=row, column=col, value=value if value is not None else "")
            cell.border = _thin_border()
            if position % 2:
                cell.fill = PatternFill("solid", fgColor=_BAND)
            color = _threshold_color(name, value)
            if color:
                cell.font = Font(color=color, bold=True)
        row += 1
    return row


def _write_facts(sheet, row: int, facts: dict[str, Any]) -> int:
    for key, value in facts.items():
        sheet.cell(row=row, column=1, value=key).font = Font(bold=True, color=_SLATE)
        cell = sheet.cell(row=row, column=2, value=value if value is not None else "")
        color = _threshold_color(key, value)
        if color:
            cell.font = Font(color=color, bold=True)
        row += 1
    return row


def _autosize(sheet, max_width: int = 55) -> None:
    widths: dict[str, int] = {}
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            widths[cell.column_letter] = max(widths.get(cell.column_letter, 0), len(str(cell.value)))
    for letter, width in widths.items():
        sheet.column_dimensions[letter].width = min(max_width, max(11, width + 2))


def _title(sheet, text: str, ncols: int, size: int = 16) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(ncols, 2))
    cell = sheet.cell(row=1, column=1, value=text)
    cell.font = Font(bold=True, size=size, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=_INK)
    cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 26


def _summary_sheet(sheet, payload: dict[str, Any], snapshots: list[DatabaseSnapshot]) -> None:
    rows = summary_rows(snapshots)
    _title(sheet, "Rapport standard du parc Oracle", len(rows[0]) if rows else 6)
    sheet.cell(row=2, column=1, value=f"Genere le : {payload.get('generated_at', '')}").font = Font(
        color=_SLATE
    )
    sheet.cell(row=3, column=1, value=f"Inventaire : {payload.get('inventory', '')}").font = Font(
        color=_SLATE
    )

    stats = kpis(snapshots)
    row = 5
    row = _table_header(sheet, row, ["Indicateur", "Valeur"])
    overview = [
        ("Bases", stats["bases"]),
        ("Collectes OK", stats["ok"]),
        ("En erreur", stats["errors"]),
        ("Taille totale (Go)", stats["total_gb"]),
        ("Tablespace le plus rempli (%)", stats["max_tablespace_pct"]),
        ("Objets invalides", stats["invalid_total"]),
        ("Comptes verrouilles", stats["locked_total"]),
        ("Sauvegarde la plus ancienne (j)", stats["oldest_backup_days"]),
    ]
    for key, value in overview:
        sheet.cell(row=row, column=1, value=key).font = Font(bold=True, color=_SLATE)
        cell = sheet.cell(row=row, column=2, value=value if value is not None else "")
        color = _threshold_color(key, value)
        if color:
            cell.font = Font(color=color, bold=True)
        row += 1

    row += 1
    header_row = row
    if rows:
        end = _table_rows(sheet, row, rows)
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(rows[0]))}{end - 1}"
    _autosize(sheet)
    sheet.freeze_panes = f"A{header_row + 1}"


def write_workbook(
    payload: dict[str, Any], snapshots: list[DatabaseSnapshot], path: Path
) -> Path:
    """Ecrit un classeur Excel : une feuille de synthese + une feuille par base."""

    if not _HAS_OPENPYXL:
        raise RuntimeError(
            "Le format Excel necessite openpyxl. Installez-le "
            "(pip install openpyxl) ou retirez 'xlsx' des formats demandes."
        )

    workbook = Workbook()
    used: set[str] = set()

    summary_sheet = workbook.active
    summary_sheet.title = _sheet_safe_name("Synthese", used)
    summary_sheet.sheet_properties.tabColor = _ACCENT
    _summary_sheet(summary_sheet, payload, snapshots)

    for snapshot in snapshots:
        ncols = 6
        for section in snapshot.sections:
            for table_rows in section.tables.values():
                if table_rows:
                    ncols = max(ncols, len(_csv_columns(table_rows)))

        sheet = workbook.create_sheet(_sheet_safe_name(snapshot.name, used))
        sheet.sheet_properties.tabColor = _OK if snapshot.ok else _ERR
        _title(sheet, snapshot.name, ncols, size=15)
        sheet.cell(
            row=2,
            column=1,
            value=(
                f"Env : {snapshot.environment or '-'} | Service : {snapshot.service or '-'} | "
                f"DSN : {snapshot.dsn} | Collecte : {snapshot.collected_at}"
            ),
        ).font = Font(color=_SLATE)

        row = 3
        if not snapshot.ok:
            sheet.cell(
                row=row, column=1, value=f"ERREUR DE COLLECTE : {snapshot.error}"
            ).font = Font(bold=True, color=_ERR)
            row += 1
        row += 1
        header_row = row

        for section in snapshot.sections:
            row = _band(sheet, row, ncols, f"{section.section}  ·  {section.status}")
            row = _write_facts(sheet, row, section.facts)
            for table, table_rows in section.tables.items():
                sheet.cell(row=row, column=1, value=table).font = Font(bold=True, color=_ACCENT)
                row += 1
                if table_rows:
                    row = _table_rows(sheet, row, table_rows)
                else:
                    sheet.cell(row=row, column=1, value="(aucune ligne)").font = Font(
                        italic=True, color=_MUTED
                    )
                    row += 1
                row += 1
            if section.messages:
                sheet.cell(row=row, column=1, value="Messages / avertissements").font = Font(
                    bold=True, color=_SLATE
                )
                row += 1
                for message in section.messages:
                    sheet.cell(row=row, column=1, value=message)
                    row += 1
            row += 1
        _autosize(sheet)
        sheet.freeze_panes = f"A{header_row}"

    workbook.save(path)
    return path


def write_reports(
    payload: dict[str, Any],
    snapshots: list[DatabaseSnapshot],
    directory: Path,
    formats: Iterable[str],
) -> list[Path]:
    """Genere les rapports demandes dans le dossier fourni."""

    directory.mkdir(parents=True, exist_ok=True)
    wanted = {item.strip().lower() for item in formats}
    if "all" in wanted:
        wanted = {"md", "html", "csv", "xlsx"}

    written: list[Path] = []
    if {"md", "markdown"} & wanted:
        path = directory / "rapport.md"
        path.write_text(render_markdown(payload, snapshots), encoding="utf-8")
        written.append(path)
    if "html" in wanted:
        path = directory / "rapport.html"
        path.write_text(render_html(payload, snapshots), encoding="utf-8")
        written.append(path)
    if "csv" in wanted:
        written.extend(write_csv(snapshots, directory))
    if "xlsx" in wanted:
        path = directory / "rapport.xlsx"
        write_workbook(payload, snapshots, path)
        written.append(path)
    return written
